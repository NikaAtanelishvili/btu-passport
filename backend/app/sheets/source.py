"""Where the three registers come from (SH-01): the Google Sheet, or the same three tabs as an
.xlsx or a folder of CSVs (tests, the seed export, Excel import).

A source returns raw cell values keyed by the canonical header names of the addendum's §4;
checking them is the sync engine's job. Headers match case-insensitively after trimming, and
`app/data/sheet_mapping.json` can rename tabs and headers for a differently laid-out copy.
"""

import csv
import io
import json
from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Protocol
from urllib.parse import quote

import httpx
from google.auth import transport
from google.oauth2 import service_account

from app.config import REPO_ROOT, get_settings

STUDENTS = "Students"
SKILLS = "Skills Catalog"
VERIFIED = "Verified Skills"
TABS = (STUDENTS, SKILLS, VERIFIED)

# (required, optional) columns per tab, in the order the export writes them.
COLUMNS: dict[str, tuple[tuple[str, ...], tuple[str, ...]]] = {
    STUDENTS: (
        ("student_id", "first_name", "last_name", "email", "program", "education_level", "status"),
        (),
    ),
    SKILLS: (("skill_code", "name", "category", "description"), ("badge_code",)),
    VERIFIED: (
        ("record_id", "student_id", "skill_code", "date_earned", "status"),
        ("earned_through",),
    ),
}
# Written by the system on the Verified Skills tab when SHEET_WRITEBACK=true (SH-06).
SYSTEM_COLUMNS = ("sync_status", "sync_message", "synced_at")
# File names of a CSV folder; also what `cli export-sheets` writes.
CSV_FILES = {
    STUDENTS: "students.csv",
    SKILLS: "skills_catalog.csv",
    VERIFIED: "verified_skills.csv",
}
MAPPING_FILE = Path(__file__).resolve().parents[1] / "data" / "sheet_mapping.json"

SHEETS_API = "https://sheets.googleapis.com/v4/spreadsheets"
SCOPE_READ = "https://www.googleapis.com/auth/spreadsheets.readonly"
SCOPE_WRITE = "https://www.googleapis.com/auth/spreadsheets"
TIMEOUT = 30.0

# A raw cell: text from CSV, typed values from .xlsx, numbers (dates as serial numbers) from
# the Sheets API's unformatted values.
Cell = str | int | float | bool | date | datetime


class SourceError(Exception):
    """The source couldn't be read (network, permission, missing tab or file). The message
    is safe to show and store: it never holds credentials."""


@dataclass(frozen=True)
class Row:
    number: int  # the row's number in the tab (1 = first row), for messages and write-back
    values: dict[str, Cell]  # canonical header -> value; "" for an empty cell


@dataclass(frozen=True)
class Tab:
    name: str  # canonical tab name (STUDENTS, SKILLS, VERIFIED)
    title: str  # the tab's title in the source
    header_row: int
    columns: dict[str, int]  # canonical header -> 0-based column index
    rows: list[Row] = field(default_factory=list)

    def missing(self) -> list[str]:
        """Required columns the tab doesn't have."""
        return [c for c in COLUMNS[self.name][0] if c not in self.columns]


class SheetSource(Protocol):
    label: str  # shown in the run summary, e.g. "Google Sheet" or "upload passports.xlsx"

    def read(self) -> dict[str, Tab]: ...


@dataclass(frozen=True)
class SheetMapping:
    titles: dict[str, str]  # canonical tab name -> title in the source
    headers: dict[str, dict[str, str]]  # canonical tab -> {source header (lower): canonical}
    header_rows: dict[str, int]


@lru_cache
def mapping() -> SheetMapping:
    """The optional mapping file; identity when absent."""
    raw: dict[str, Any] = {}
    if MAPPING_FILE.exists():
        raw = json.loads(MAPPING_FILE.read_text(encoding="utf-8"))
    tabs: dict[str, Any] = raw.get("tabs", {})
    return SheetMapping(
        titles={t: str(tabs.get(t, {}).get("title", t)) for t in TABS},
        headers={
            t: {k.strip().lower(): v for k, v in tabs.get(t, {}).get("headers", {}).items()}
            for t in TABS
        },
        header_rows={t: int(tabs.get(t, {}).get("header_row", 1)) for t in TABS},
    )


def _canonical(tab: str, header: object) -> str:
    key = str(header).strip().lower()
    return mapping().headers[tab].get(key, key)


def build_tab(name: str, title: str, grid: list[list[Cell | None]]) -> Tab:
    """Turn a tab's cells (row lists, as read) into header-keyed rows. Blank rows are skipped;
    a later column with a header already seen is ignored."""
    header_row = mapping().header_rows[name]
    if len(grid) < header_row:
        return Tab(name=name, title=title, header_row=header_row, columns={})
    columns: dict[str, int] = {}
    for i, header in enumerate(grid[header_row - 1]):
        if header is not None and str(header).strip():
            columns.setdefault(_canonical(name, header), i)
    rows = []
    for offset, cells in enumerate(grid[header_row:], start=header_row + 1):
        values: dict[str, Cell] = {}
        for column, i in columns.items():
            value = cells[i] if i < len(cells) else None
            values[column] = "" if value is None else value
        if any(str(v).strip() for v in values.values()):
            rows.append(Row(number=offset, values=values))
    return Tab(name=name, title=title, header_row=header_row, columns=columns, rows=rows)


# --- Files -----------------------------------------------------------------------------------


class FileSource:
    """An .xlsx with the three tabs (by title), or a folder holding the three CSV_FILES."""

    def __init__(self, path: Path | None = None, *, xlsx: bytes | None = None, label: str = ""):
        if (path is None) == (xlsx is None):
            raise ValueError("give a path or xlsx bytes")
        self.path = path
        self.xlsx = xlsx
        self.label = label or (f"file {path.name}" if path else "upload")

    def read(self) -> dict[str, Tab]:
        if self.xlsx is not None:
            return _read_xlsx(io.BytesIO(self.xlsx))
        assert self.path is not None
        if self.path.is_dir():
            return _read_csv_folder(self.path)
        if not self.path.exists():
            raise SourceError(f"{self.path.name} not found")
        return _read_xlsx(self.path)


def _read_csv_folder(folder: Path) -> dict[str, Tab]:
    tabs = {}
    for name in TABS:
        path = folder / CSV_FILES[name]
        if not path.exists():
            raise SourceError(f"Tab {name!r} is missing: no {path.name} in {folder.name}/")
        with path.open(newline="", encoding="utf-8-sig") as fh:
            grid: list[list[Cell | None]] = [list(r) for r in csv.reader(fh)]
        tabs[name] = build_tab(name, path.name, grid)
    return tabs


def _read_xlsx(file: Path | io.BytesIO) -> dict[str, Tab]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:  # zip or XML errors: not an .xlsx
        raise SourceError(f"Not a readable .xlsx file ({type(exc).__name__})") from None
    titles = {t.strip().lower(): t for t in book.sheetnames}
    tabs = {}
    for name in TABS:
        title = titles.get(mapping().titles[name].strip().lower())
        if title is None:
            raise SourceError(f"Tab {mapping().titles[name]!r} is missing from the workbook")
        grid = [[_xlsx_cell(v) for v in r] for r in book[title].iter_rows(values_only=True)]
        tabs[name] = build_tab(name, title, grid)
    book.close()
    return tabs


def _xlsx_cell(value: object) -> Cell | None:
    if value is None or isinstance(value, str | int | float | bool | date | datetime):
        return value
    return str(value)  # times, decimals, rich text: as text


# --- Google Sheets ---------------------------------------------------------------------------


class _HttpxResponse(transport.Response):
    def __init__(self, response: httpx.Response) -> None:
        self._r = response

    @property
    def status(self) -> int:
        return self._r.status_code

    @property
    def headers(self) -> Mapping[str, str]:
        return self._r.headers

    @property
    def data(self) -> bytes:
        return self._r.content


class _HttpxRequest(transport.Request):
    """google-auth's transport interface over httpx (it ships one for `requests` only)."""

    def __init__(self, client: httpx.Client) -> None:
        self.client = client

    def __call__(
        self,
        url: str,
        method: str = "GET",
        body: bytes | None = None,
        headers: Mapping[str, str] | None = None,
        timeout: float | None = None,
        **kwargs: Any,
    ) -> _HttpxResponse:
        try:
            response = self.client.request(
                method, url, content=body, headers=headers, timeout=timeout or TIMEOUT
            )
        except httpx.HTTPError as exc:
            from google.auth import exceptions

            error = exceptions.TransportError(type(exc).__name__)  # type: ignore[no-untyped-call]
            raise error from None
        return _HttpxResponse(response)


def _key_info() -> dict[str, Any]:
    s = get_settings()
    content = s.google_sheets_credentials_json.get_secret_value().strip()
    if content:
        try:
            info: dict[str, Any] = json.loads(content)
        except ValueError:
            raise SourceError("GOOGLE_SHEETS_CREDENTIALS_JSON is not valid JSON") from None
        return info
    if not s.google_sheets_credentials_file:
        raise SourceError(
            "No service account key: set GOOGLE_SHEETS_CREDENTIALS_FILE or _JSON in .env"
        )
    path = Path(s.google_sheets_credentials_file)
    path = path if path.is_absolute() else REPO_ROOT / path
    try:
        info = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise SourceError(f"Can't read the service account key file {path.name}") from None
    return info


@lru_cache
def _credentials(scope: str) -> service_account.Credentials:
    """One credentials object per process and scope: its access token lasts an hour."""
    try:
        info = _key_info()
        creds: service_account.Credentials = service_account.Credentials.from_service_account_info(  # type: ignore[no-untyped-call]
            info, scopes=[scope]
        )
    except (KeyError, ValueError):
        raise SourceError("The service account key is incomplete or malformed") from None
    return creds


def _a1_tab(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


class GoogleSheetSource:
    """The Sheet named by GOOGLE_SHEETS_ID, read with one values.batchGet. The scope is
    read-only unless SHEET_WRITEBACK=true."""

    label = "Google Sheet"

    def __init__(self, sheet_id: str | None = None, *, writeback: bool | None = None) -> None:
        s = get_settings()
        self.sheet_id = sheet_id if sheet_id is not None else s.google_sheets_id
        if not self.sheet_id:
            raise SourceError("GOOGLE_SHEETS_ID is not set")
        self.writeback = s.sheet_writeback if writeback is None else writeback
        self.client = httpx.Client(timeout=TIMEOUT)

    def _token(self) -> str:
        from google.auth import exceptions

        creds = _credentials(SCOPE_WRITE if self.writeback else SCOPE_READ)
        if not creds.valid:
            try:
                creds.refresh(_HttpxRequest(self.client))  # type: ignore[no-untyped-call]
            except exceptions.GoogleAuthError as exc:
                # The message is Google's error code (e.g. invalid_grant), never the key.
                raise SourceError(f"Google sign-in for the service account failed: {exc}") from None
        return str(creds.token)

    def _call(self, method: str, url: str, **kwargs: Any) -> dict[str, Any]:
        headers = {"Authorization": f"Bearer {self._token()}"}
        try:
            response = self.client.request(method, url, headers=headers, **kwargs)
        except httpx.HTTPError as exc:
            raise SourceError(f"Google Sheets unreachable ({type(exc).__name__})") from None
        if response.status_code != 200:
            raise SourceError(_api_error(response))
        body: dict[str, Any] = response.json()
        return body

    def read(self) -> dict[str, Tab]:
        titles = mapping().titles
        body = self._call(
            "GET",
            f"{SHEETS_API}/{quote(self.sheet_id)}/values:batchGet",
            params=[
                *[("ranges", _a1_tab(titles[name])) for name in TABS],
                ("valueRenderOption", "UNFORMATTED_VALUE"),
                ("dateTimeRenderOption", "SERIAL_NUMBER"),
                ("majorDimension", "ROWS"),
            ],
        )
        ranges = body.get("valueRanges", [])
        if len(ranges) != len(TABS):
            raise SourceError("Google Sheets returned an unexpected answer")
        return {
            name: build_tab(name, titles[name], [list(r) for r in vr.get("values", [])])
            for name, vr in zip(TABS, ranges, strict=True)
        }

    def write(self, data: list[tuple[str, list[list[str]]]]) -> int:
        """values.batchUpdate of (A1 range, rows) pairs, as plain text. Returns cells written."""
        if not data:
            return 0
        body = self._call(
            "POST",
            f"{SHEETS_API}/{quote(self.sheet_id)}/values:batchUpdate",
            json={
                "valueInputOption": "RAW",
                "data": [{"range": r, "values": v} for r, v in data],
            },
        )
        return int(body.get("totalUpdatedCells", 0))


def _api_error(response: httpx.Response) -> str:
    try:
        error = response.json().get("error", {})
        status, message = error.get("status", ""), error.get("message", "")
    except ValueError:
        status, message = "", ""
    if response.status_code == 403:
        return (
            "The service account has no access to the Sheet: share it with the service "
            f"account's email ({status or 'PERMISSION_DENIED'})"
        )
    if response.status_code == 404:
        return "No Sheet with this GOOGLE_SHEETS_ID"
    if response.status_code == 400 and "parse range" in message:
        expected = ", ".join(repr(mapping().titles[t]) for t in TABS)
        return f"A tab is missing: the Sheet needs tabs named exactly {expected}"
    return f"Google Sheets error {response.status_code} {status}: {message[:200]}".strip()
