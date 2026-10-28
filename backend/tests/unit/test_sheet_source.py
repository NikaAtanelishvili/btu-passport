"""SH-01: the three tabs from an .xlsx, a CSV folder and the Sheets API (faked)."""

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
import pytest
from openpyxl import Workbook

from app.sheets import source
from app.sheets.source import SKILLS, STUDENTS, VERIFIED, FileSource, GoogleSheetSource

STUDENT_HEADER = [
    "student_id", "first_name", "last_name", "email", "program", "education_level", "status",
]  # fmt: skip
STUDENT = ["BTU-T-0001", "Giorgi", "Kapanadze", "btu-t-0001@example.org", "IT", "BACHELOR",
           "ACTIVE"]  # fmt: skip
SKILL_HEADER = ["skill_code", "name", "category", "description", "badge_code"]
SKILL = ["SKL-001", "AI Literacy", "AI Skill", "Fundamentals.", "BDG-AI-001"]
VERIFIED_HEADER = ["record_id", "student_id", "skill_code", "date_earned", "status"]


def write_xlsx(path: Path, verified_date: object) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = STUDENTS
    sheet.append([f"  {h.upper()} " for h in STUDENT_HEADER])  # case and spaces don't matter
    sheet.append(STUDENT)
    sheet.append([None] * 7)  # blank row: skipped
    sheet.append(["BTU-T-0002", "Mariam", "Beridze", "b@example.org", "IT", "MASTER", "ACTIVE"])
    skills = book.create_sheet(SKILLS)
    skills.append([*SKILL_HEADER, "notes"])  # extra column: kept, ignored by the engine
    skills.append([*SKILL, "x"])
    verified = book.create_sheet(VERIFIED)
    verified.append(VERIFIED_HEADER)
    verified.append(["REC-0001", "BTU-T-0001", "SKL-001", verified_date, "CONFIRMED"])
    book.save(path)
    return path


def test_xlsx(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "data.xlsx", datetime(2026, 9, 18))
    tabs = FileSource(path).read()
    assert list(tabs) == [STUDENTS, SKILLS, VERIFIED]
    students = tabs[STUDENTS]
    assert [r.number for r in students.rows] == [2, 4]
    assert students.rows[0].values == dict(zip(STUDENT_HEADER, STUDENT, strict=True))
    assert students.missing() == []
    assert tabs[SKILLS].rows[0].values["notes"] == "x"
    assert tabs[VERIFIED].rows[0].values["date_earned"] == datetime(2026, 9, 18)


def test_xlsx_bytes_and_missing_tab(tmp_path: Path) -> None:
    data = write_xlsx(tmp_path / "data.xlsx", "2026-09-18").read_bytes()
    assert len(FileSource(xlsx=data).read()[VERIFIED].rows) == 1
    book = Workbook()
    assert book.active is not None
    book.active.title = STUDENTS
    book.save(tmp_path / "one-tab.xlsx")
    with pytest.raises(source.SourceError, match="'Skills Catalog' is missing"):
        FileSource(tmp_path / "one-tab.xlsx").read()
    with pytest.raises(source.SourceError, match=r"Not a readable \.xlsx"):
        FileSource(xlsx=b"not a zip").read()


def test_csv_folder(tmp_path: Path) -> None:
    (tmp_path / "students.csv").write_text(
        ",".join(STUDENT_HEADER[:-1]) + "\n" + ",".join(STUDENT[:-1]) + "\n,,,,,,\n"
    )
    (tmp_path / "skills_catalog.csv").write_text(
        ",".join(SKILL_HEADER) + '\nSKL-001,AI Literacy,AI Skill,"Uses, limits.",\n'
    )
    (tmp_path / "verified_skills.csv").write_text(",".join(VERIFIED_HEADER) + "\n")
    tabs = FileSource(tmp_path).read()
    assert len(tabs[STUDENTS].rows) == 1
    assert tabs[STUDENTS].missing() == ["status"]  # the engine fails this tab
    assert tabs[SKILLS].rows[0].values["description"] == "Uses, limits."
    assert tabs[SKILLS].rows[0].values["badge_code"] == ""
    assert tabs[VERIFIED].rows == []
    (tmp_path / "students.csv").unlink()
    with pytest.raises(source.SourceError, match=r"no students\.csv"):
        FileSource(tmp_path).read()


def test_mapping_renames_tabs_and_headers(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    mapping = {
        "tabs": {
            STUDENTS: {
                "title": "Students (test)",
                "header_row": 2,
                "headers": {"Student ID": "student_id", "Given name": "first_name"},
            }
        }
    }
    file = tmp_path / "mapping.json"
    file.write_text(json.dumps(mapping))
    monkeypatch.setattr(source, "MAPPING_FILE", file)
    source.mapping.cache_clear()
    try:
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "STUDENTS (TEST)"
        sheet.append(["BTU test data"])
        sheet.append(["Student ID", "Given name"])
        sheet.append(["BTU-T-0001", "Giorgi"])
        book.create_sheet(SKILLS)
        book.create_sheet(VERIFIED)
        book.save(tmp_path / "x.xlsx")
        tab = FileSource(tmp_path / "x.xlsx").read()[STUDENTS]
        assert tab.title == "STUDENTS (TEST)"
        assert tab.rows[0].number == 3
        assert tab.rows[0].values == {"student_id": "BTU-T-0001", "first_name": "Giorgi"}
    finally:
        source.mapping.cache_clear()


class FakeCreds:
    valid = True
    token = "test-token"


def google(monkeypatch: pytest.MonkeyPatch, handler: Any) -> GoogleSheetSource:
    monkeypatch.setattr(source, "_credentials", lambda scope: FakeCreds())
    sheet = GoogleSheetSource("sheet-123", writeback=False)
    sheet.client = httpx.Client(transport=httpx.MockTransport(handler))
    return sheet


def test_google_reads_three_tabs_in_one_call(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(
            200,
            json={
                "valueRanges": [
                    {"values": [STUDENT_HEADER, STUDENT]},
                    {"values": [SKILL_HEADER]},
                    {"values": [VERIFIED_HEADER, ["REC-0001", "BTU-T-0001", "SKL-001", 46283]]},
                ]
            },
        )

    tabs = google(monkeypatch, handler).read()
    assert len(calls) == 1
    request = calls[0]
    assert request.url.path == "/v4/spreadsheets/sheet-123/values:batchGet"
    assert request.url.params.get_list("ranges") == [
        "'Students'",
        "'Skills Catalog'",
        "'Verified Skills'",
    ]
    assert request.url.params["valueRenderOption"] == "UNFORMATTED_VALUE"
    assert request.headers["Authorization"] == "Bearer test-token"
    assert tabs[STUDENTS].rows[0].values["student_id"] == "BTU-T-0001"
    assert tabs[SKILLS].rows == []
    row = tabs[VERIFIED].rows[0]
    assert row.values["date_earned"] == 46283  # a serial date; the engine converts it
    assert row.values["status"] == ""  # short row: missing trailing cells are empty


@pytest.mark.parametrize(
    ("status", "body", "message"),
    [
        (403, {"error": {"status": "PERMISSION_DENIED"}}, "share it with the service account"),
        (404, {"error": {"status": "NOT_FOUND"}}, "No Sheet with this GOOGLE_SHEETS_ID"),
        (
            400,
            {"error": {"message": "Unable to parse range: 'Students'"}},
            "tabs named exactly 'Students'",
        ),
    ],
)
def test_google_errors(
    monkeypatch: pytest.MonkeyPatch, status: int, body: dict[str, Any], message: str
) -> None:
    sheet = google(monkeypatch, lambda request: httpx.Response(status, json=body))
    with pytest.raises(source.SourceError, match=message):
        sheet.read()


def test_google_unreachable(monkeypatch: pytest.MonkeyPatch) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("no route")

    with pytest.raises(source.SourceError, match="unreachable"):
        google(monkeypatch, handler).read()


def test_no_sheet_id_or_key(monkeypatch: pytest.MonkeyPatch) -> None:
    with pytest.raises(source.SourceError, match="GOOGLE_SHEETS_ID is not set"):
        GoogleSheetSource("")
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "google_sheets_credentials_file", "missing.json")
    with pytest.raises(source.SourceError, match="Can't read the service account key"):
        source._key_info()


def test_key_from_json_content(monkeypatch: pytest.MonkeyPatch) -> None:
    """The demo host passes the key's content (GOOGLE_SHEETS_CREDENTIALS_JSON), not a file."""
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa

    from app.config import get_settings

    pem = (
        rsa.generate_private_key(public_exponent=65537, key_size=2048)
        .private_bytes(
            serialization.Encoding.PEM,
            serialization.PrivateFormat.PKCS8,
            serialization.NoEncryption(),
        )
        .decode()
    )
    info = {
        "type": "service_account",
        "client_email": "sync@test.iam.gserviceaccount.com",
        "private_key": pem,
        "private_key_id": "1",
        "token_uri": "https://oauth2.googleapis.com/token",
    }
    from pydantic import SecretStr

    monkeypatch.setattr(
        get_settings(), "google_sheets_credentials_json", SecretStr(json.dumps(info))
    )
    source._credentials.cache_clear()
    try:
        creds = source._credentials(source.SCOPE_READ)
        assert creds.service_account_email == "sync@test.iam.gserviceaccount.com"
        assert creds.scopes == [source.SCOPE_READ]
    finally:
        source._credentials.cache_clear()
    monkeypatch.setattr(get_settings(), "google_sheets_credentials_json", SecretStr("{nope"))
    with pytest.raises(source.SourceError, match="not valid JSON"):
        source._key_info()
